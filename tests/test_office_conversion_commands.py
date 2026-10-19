from __future__ import annotations

from io import BytesIO
import json
import queue
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pytest

from jarvis.llm import AIError, GeminiAI, OpenAICompatibleAI, _json_object
from jarvis.office.commands import local_conversion_intent
from jarvis.office.converters import convert_file
from jarvis.office.service import OfficeService, validate_request
from jarvis.office.destinations import wants_desktop
from jarvis.ui.app import JarvisApp
from jarvis.intent import AIIntent


@pytest.mark.parametrize("path", [r"C:\Users\WD\Documents\ChatGPT\jarvis\data\office\outputs\ozbekiston_taqdimot_9192df5c24.pptx", r"C:\My Files\O'zbekiston.pptx", "/tmp/my slides.pptx"])
def test_explicit_pptx_word_command_is_local_and_path_is_exact(path):
    intent = local_conversion_intent(path + " ushbu faylni word fayl qilib ber")
    assert intent.is_command and intent.command == "office_convert"
    assert json.loads(intent.target) == {"source": path, "target_format": "docx"}
    assert json.loads(intent.target)["source"] == path


def test_local_conversion_never_calls_ai():
    app = SimpleNamespace(ai=Mock(), events=queue.Queue())
    JarvisApp._reply(app, r"C:\Folder\slides.pptx ushbu faylni word fayl qilib ber")
    app.ai.interpret.assert_not_called()
    kind, intent = app.events.get_nowait()
    assert kind == "intent" and intent.command == "office_convert"


@pytest.mark.parametrize("message", [
    "Wordga aylantirib ber", "https://example.com/slides.pptx Wordga aylantir",
    r"C:\a.pptx C:\b.pptx Wordga aylantir", r"C:\a.pptx ni qanday qilib Wordga aylantirish mumkinmi",
    r"C:\a.pptx qayerda", r"C:\a.pptx Word yoki PDF qil",
])
def test_ambiguous_questions_or_remote_sources_are_not_guessed(message):
    assert local_conversion_intent(message) is None


def test_pptx_word_preserves_text_tables_unicode_and_source(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document
    source = tmp_path / "Ozbekiston.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text_frame.text = "O'zbekiston — Vatanim"
    table = slide.shapes.add_table(2, 2, 0, Inches(2), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "Shahar"
    table.cell(1, 0).text = "Toshkent"
    table.cell(1, 1).text = "Poytaxt"
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(5), Inches(1)).text_frame.text = "Xulosa"
    deck.save(source)
    original = source.read_bytes()
    target = tmp_path / "converted.docx"
    result = convert_file(source, target)
    document = Document(target)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Vatanim" in text and "Xulosa" in text and "Slayd 2" in text
    assert document.tables[0].cell(1, 0).text == "Toshkent"
    assert "slayd bezagi" in result and source.read_bytes() == original
    with pytest.raises(ValueError, match="mavjud"):
        convert_file(source, target)


def test_empty_pptx_does_not_produce_a_misleading_word_document(tmp_path):
    from pptx import Presentation
    source, target = tmp_path / "empty.pptx", tmp_path / "empty.docx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    with pytest.raises(ValueError, match="matn/jadval"):
        convert_file(source, target)
    assert not target.exists()


def test_fenced_router_json_and_multipart_gemini_response_are_supported():
    content = json.dumps({"type": "command", "reply": "Lokal aylantirish", "category": "office", "command": "office_convert", "target": {"source": "C:/docs/source.pptx", "target_format": "docx"}})
    midpoint = len(content)//2
    data = {"candidates": [{"finishReason": "STOP", "content": {"parts": [
        {"thought": True, "text": "Not a JSON answer"}, {"text": "```json\n"+content[:midpoint]}, {"text": content[midpoint:]+"\n```"}]}}]}
    with patch("jarvis.llm.urlopen", return_value=BytesIO(json.dumps(data).encode())):
        intent = GeminiAI("fake", "fake").interpret("Convert source")
    assert intent.command == "office_convert"
    assert json.loads(intent.target)["source"] == "C:/docs/source.pptx"


def test_router_rejects_arrays_and_does_not_guess_invalid_backslashes():
    for text in ("[]", "null", r'{"source":"C:\Users\new\slides.pptx"}'):
        with pytest.raises(ValueError):
            _json_object(text)
    assert _json_object('```json\n{"reply":"Hello"}\n```') == {"reply": "Hello"}


@pytest.mark.parametrize("message,expected", [("desktopda saqla", True), ("ish stoliga saqla", True),
    ("save to desktop", True), ("desktopda saqlama", False), ("desktop nima", False),
    (r"C:\folder\desktopda saqla report.pptx Wordga aylantir", False)])
def test_desktop_destination_requires_an_explicit_positive_request(message, expected):
    assert wants_desktop(message) is expected


def test_desktop_copy_never_overwrites_existing_file(tmp_path):
    output, desktop = tmp_path / "outputs", tmp_path / "desktop"
    desktop.mkdir()
    source = tmp_path / "source.txt"
    source.write_text("Original", encoding="utf-8")
    old = desktop / "fixed.docx"
    old.write_bytes(b"Keep this")
    service = OfficeService(output)
    with patch("jarvis.office.service.desktop_directory", return_value=desktop), patch.object(service, "_filename", side_effect=["fixed.docx", "unique.docx"]):
        result = service.run("office_convert", {"source": str(source), "target_format": "docx", "save_to": "desktop"})
    assert result.extra_paths[0].parent == desktop
    assert result.extra_paths[0].read_bytes() == result.path.read_bytes()
    assert old.read_bytes() == b"Keep this"
    assert "Desktop nusxasi" in result.summary


def test_output_destination_cannot_be_an_arbitrary_model_path():
    with pytest.raises(ValueError, match="library yoki desktop"):
        validate_request("office_create", {"format": "docx", "prompt": "Make a document", "save_to": "C:/private"})
    intent = AIIntent("command", "Okay", "office", "office_create", json.dumps({"format": "docx", "prompt": "Make report", "save_to": "desktop"}))
    assert "save_to" not in json.loads(JarvisApp._validate_office_intent("Hisobot tayyorla", intent).target)


def test_creation_honors_desktop_request_in_original_prompt(tmp_path):
    desktop = tmp_path / "desktop"
    desktop.mkdir()
    ai = Mock()
    ai.complete_json.return_value = {"title": "Ona Vatan", "sections": [{"paragraphs": ["Test mazmuni"]}]}
    with patch("jarvis.office.service.desktop_directory", return_value=desktop):
        result = OfficeService(tmp_path / "outputs").run("office_create", {"format": "docx", "prompt": "Ona vatan haqida esse yoz va desktopda saqla"}, ai=ai)
    assert len(result.extra_paths) == 1
    assert result.extra_paths[0].is_file()
    assert result.extra_paths[0].parent == desktop
