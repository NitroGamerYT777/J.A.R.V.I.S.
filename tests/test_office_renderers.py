from __future__ import annotations

from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import pytest

from jarvis.office.renderers import render_document


def word_spec():
    return {"title": "2026 ish rejasi", "subtitle": "Tahrirlanadigan loyiha", "sections": [
        {"heading": "Maqsad", "paragraphs": ["Jamoa uchun ish rejasini tayyorlash."], "bullets": ["Yangi vazifalar", "Natijani tekshirish"],
         "table": {"headers": ["Vazifa", "Soat"], "rows": [["Tahlil", 4], ["Sinov", 2]]}}]}


def excel_spec():
    return {"title": "Oylik hisob", "sheets": [{"name": "Hisob", "headers": ["Nomi", "Soni", "Narxi", "Jami"],
        "rows": [["Daftar", 2, 12000, None], ["Qalam", 3, 5000, None], ["Jami", None, None, None]],
        "formulas": [{"cell": "D2", "formula": "=B2*C2"}, {"cell": "D3", "formula": "=B3*C3"}, {"cell": "D4", "formula": "=SUM(D2:D3)"}],
        "chart": {"type": "bar", "title": "Mahsulotlar", "category_column": 1, "value_columns": [2]}}]}


def ppt_spec():
    return {"title": "Loyiha taqdimoti", "slides": [
        {"title": "Loyiha maqsadi", "body": "O‘zbekcha ishlaydigan yordamchi", "bullets": ["Hujjatlar tayyorlash", "Vazifalarni boshqarish"], "notes": "Bu tahrirlanadigan slayd."},
        {"title": "Ish rejasi", "table": {"headers": ["Bosqich", "Muddat"], "rows": [["Tahlil", "1 kun"], ["Tayyorlash", "2 kun"]]}}]}


def test_docx_contains_editable_text_lists_and_styled_table(tmp_path):
    target = tmp_path / "report.docx"
    render_document(word_spec(), target)
    doc = Document(target)
    assert doc.paragraphs[0].text == "2026 ish rejasi"
    assert doc.paragraphs[0].style.name == "Title"
    assert str(doc.styles["Title"].font.color.rgb) == "000000"
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert doc.tables[0].cell(1, 0).text == "Tahlil"
    assert "tblHeader" in doc.tables[0].rows[0]._tr.xml
    assert "203B5F" in doc.tables[0].cell(0, 0)._tc.xml
    assert "D9D9D9" in doc.tables[0]._tbl.xml
    with ZipFile(target) as archive:
        assert not any("vba" in name.lower() or "embeddings/" in name for name in archive.namelist())


def test_xlsx_formulas_editable_chart_styles_and_recalculation(tmp_path):
    target = tmp_path / "budget.xlsx"
    render_document(excel_spec(), target)
    book = load_workbook(target)
    sheet = book["Hisob"]
    assert sheet["D4"].value == "=SUM(D2:D3)"
    assert sheet["D4"].data_type == "f"
    assert sheet["B2"].value == 2
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].fill.fgColor.rgb.endswith("203B5F")
    assert sheet.tables and sheet.auto_filter.ref == "A1:D4"
    assert len(sheet._charts) == 1
    assert book.calculation.fullCalcOnLoad
    assert book.calculation.forceFullCalc


@pytest.mark.parametrize("text", ["=HYPERLINK(\"https://evil.invalid\",\"Click\")", "+cmd|' /C calc'!A0", "-1+2", "@SUM(A1:A2)"])
def test_xlsx_untrusted_string_is_literal_not_formula(tmp_path, text):
    spec = excel_spec()
    spec["sheets"][0]["rows"][0][0] = text
    path = tmp_path / "literal.xlsx"
    render_document(spec, path)
    sheet = load_workbook(path)["Hisob"]
    assert sheet["A2"].value == text
    assert sheet["A2"].data_type == "s"
    assert not sheet["A2"].hyperlink


def test_xlsx_local_sheet_reference_is_supported(tmp_path):
    spec = excel_spec()
    spec["sheets"].append({"name": "Jami hisob", "headers": ["Jami"], "rows": [[None]], "formulas": [{"cell": "A2", "formula": "=SUM(Hisob!D2:D3)"}]})
    spec["sheets"][0]["formulas"][0]["formula"] = "=IF('Jami hisob'!A2>0,B2*C2,0)"
    path = tmp_path / "cross-sheet.xlsx"
    render_document(spec, path)
    assert load_workbook(path)["Jami hisob"]["A2"].data_type == "f"


def test_pptx_has_native_editable_shapes_tables_and_notes(tmp_path):
    path = tmp_path / "deck.pptx"
    render_document(ppt_spec(), path)
    deck = Presentation(path)
    assert len(deck.slides) == 2
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, rel=.001)
    texts = [s.text for s in deck.slides[0].shapes if s.has_text_frame]
    assert "Loyiha maqsadi" in texts
    assert any("Hujjatlar tayyorlash" in t for t in texts)
    assert "Bu tahrirlanadigan slayd." in deck.slides[0].notes_slide.notes_text_frame.text
    table = next(shape.table for shape in deck.slides[1].shapes if shape.has_table)
    assert table.cell(1, 0).text == "Tahlil"
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "203B5F"
    assert not any(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes)


@pytest.mark.parametrize("formula", [
    '=HYPERLINK("https://evil.invalid","Open")', '=WEBSERVICE("https://evil.invalid")',
    '=cmd|" /C calc"!A1', '=[book.xlsx]Sheet1!A1', '=SUM(Unknown!A2)', '=SUM(A1:XFD1048576)',
    '=INDIRECT("A2")', '=EXEC("calc")', '=SUM(A2:A3', '=SUM(A2:A3);1', '=LAMBDA(x,x)(1)',
    '=1+', '=A2 A3', '=SUM(A2,,A3)', '=()', '=1(2)',
])
def test_unsafe_or_malformed_formula_rejected_without_partial_file(tmp_path, formula):
    spec = excel_spec()
    spec["sheets"][0]["formulas"][0]["formula"] = formula
    path = tmp_path / "unsafe.xlsx"
    with pytest.raises(ValueError):
        render_document(spec, path)
    assert not path.exists()


@pytest.mark.parametrize(("extension", "spec"), [
    ("docx", {"title": "Hisobot", "sections": "bad"}),
    ("docx", {"title": "Hisobot", "sections": [{"paragraphs": ["bad\x00text"]}]}),
    ("docx", {"title": "Hisobot", "sections": [{"table": {"headers": ["A"], "rows": [[1, 2]]}}]}),
    ("xlsx", {"title": "Hisob", "sheets": [{"name": "Bad/name", "headers": ["A"], "rows": [[1]]}]}),
    ("xlsx", {"title": "Hisob", "sheets": [{"name": "A", "headers": ["A", "a"], "rows": [[1, 2]]}]}),
    ("xlsx", {"title": "Hisob", "sheets": [{"name": "A", "headers": ["A"], "rows": [[float("inf")]]}]}),
    ("pptx", {"title": "Deck", "slides": [{"title": "A", "bullets": ["Long " * 100]}]}),
    ("pptx", {"title": "Deck", "slides": [{"title": "A", "image": "https://example.com/image.png"}]}),
])
def test_invalid_spec_rejected_before_writing(tmp_path, extension, spec):
    path = tmp_path / f"invalid.{extension}"
    with pytest.raises(ValueError):
        render_document(spec, path)
    assert not path.exists()


@pytest.mark.parametrize(("extension", "factory"), [("docx", word_spec), ("xlsx", excel_spec), ("pptx", ppt_spec)])
def test_existing_destination_never_overwritten(tmp_path, extension, factory):
    path = tmp_path / f"existing.{extension}"
    render_document(factory(), path)
    original = path.read_bytes()
    with pytest.raises(ValueError, match="mavjud"):
        render_document(factory(), path)
    assert path.read_bytes() == original


def test_sheet_cell_limit_rejects_small_strings_with_huge_shape(tmp_path):
    spec = {"title": "Big", "sheets": [{"name": "Big", "headers": [f"C{i}" for i in range(100)], "rows": [[None] * 100 for _ in range(1000)]}]}
    with pytest.raises(ValueError, match="katak"):
        render_document(spec, tmp_path / "big.xlsx")


def test_header_formula_cannot_overwrite_title(tmp_path):
    spec = excel_spec()
    spec["sheets"][0]["formulas"] = [{"cell": "A1", "formula": "=1+2"}]
    with pytest.raises(ValueError, match="sarlavha"):
        render_document(spec, tmp_path / "header.xlsx")


def test_chart_wrong_type_has_friendly_error_not_type_error(tmp_path):
    spec = excel_spec()
    spec["sheets"][0]["chart"]["type"] = []
    with pytest.raises(ValueError, match="Diagramma"):
        render_document(spec, tmp_path / "chart.xlsx")


@pytest.mark.parametrize("text", ["Bad\ufffevalue", "Bad\ud800value"])
def test_illegal_xml_characters_are_rejected_before_output(tmp_path, text):
    spec = word_spec()
    spec["sections"][0]["paragraphs"] = [text]
    path = tmp_path / "invalid-xml.docx"
    with pytest.raises(ValueError):
        render_document(spec, path)
    assert not path.exists()
