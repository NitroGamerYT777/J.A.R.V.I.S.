from __future__ import annotations

from contextlib import contextmanager
import csv
import io
import json
from pathlib import Path
import subprocess
from types import SimpleNamespace
from unittest.mock import MagicMock
import zipfile

import pytest

from jarvis.office import converters
from jarvis.office.convert_worker import _run_com
from jarvis.office import convert_worker


def test_supported_conversions_is_not_mutable_shared_state():
    first = converters.supported_conversions()
    assert first[".docx"] == (".txt", ".pdf")
    assert first[".csv"] == (".xlsx",)
    first.clear()
    assert converters.supported_conversions()


@pytest.mark.parametrize("existing", ["same", "different"])
def test_never_overwrites_source_or_existing_output(tmp_path, existing):
    source = tmp_path / "input.txt"
    source.write_text("Original o‘zbekcha", encoding="utf-8")
    destination = source if existing == "same" else tmp_path / "already.docx"
    if existing != "same":
        destination.write_bytes(b"existing result")
    original = destination.read_bytes()
    with pytest.raises(ValueError, match="ustiga|mavjud"):
        converters.convert_file(source, destination)
    assert destination.read_bytes() == original


def test_csv_xlsx_roundtrip_preserves_unicode_and_neutralizes_formulas(tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "input.csv"
    source.write_text('Ism,Qiymat\nO‘zbekiston,=1+2\nSardor," +CMD()"\nMadina,@SUM(A1)\n', encoding="utf-8-sig")
    xlsx = tmp_path / "table.xlsx"
    converters.convert_file(source, xlsx)
    workbook = load_workbook(xlsx, data_only=False)
    try:
        assert workbook.active["A2"].value == "O‘zbekiston"
        assert workbook.active["B2"].value == "'=1+2"
        assert workbook.active["B2"].data_type == "s"
        assert workbook.active["B3"].value == "' +CMD()"
    finally:
        workbook.close()
    output = tmp_path / "result.csv"
    converters.convert_file(xlsx, output)
    rows = list(csv.reader(io.StringIO(output.read_text(encoding="utf-8"))))
    assert rows[1] == ["O‘zbekiston", "'=1+2"]
    assert rows[3] == ["Madina", "'@SUM(A1)"]


def test_semicolon_csv_is_supported(tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "input.csv"
    source.write_text("Ism;Shahar\nAziz;Toshkent\n", encoding="utf-8")
    output = tmp_path / "table.xlsx"
    converters.convert_file(source, output)
    workbook = load_workbook(output)
    try:
        assert workbook.active["B2"].value == "Toshkent"
    finally:
        workbook.close()


def test_xlsx_multiple_sheets_requires_explicit_choice(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Birinchi"
    workbook.active.append(["not this"])
    workbook.create_sheet("Ikkinchi").append(["kerakli", 42])
    source = tmp_path / "table.xlsx"
    workbook.save(source)
    workbook.close()
    output = tmp_path / "chosen.csv"
    with pytest.raises(ValueError, match="Birinchi.*Ikkinchi"):
        converters.convert_file(source, output)
    assert not output.exists()
    converters.convert_file(source, output, sheet="Ikkinchi")
    assert output.read_text(encoding="utf-8").strip() == "kerakli,42"


def test_xlsx_missing_sheet_is_clear(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "table.xlsx"
    workbook = Workbook()
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="varag‘i yo‘q"):
        converters.convert_file(source, tmp_path / "out.csv", sheet="not found")


def test_formula_without_cache_is_not_silently_dropped(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "table.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+2"
    workbook.save(source)
    workbook.close()
    output = tmp_path / "out.csv"
    with pytest.raises(ValueError, match="A1.*hisoblangan qiymati yo‘q"):
        converters.convert_file(source, output)
    assert not output.exists()


def test_formula_cached_value_only_is_exported(tmp_path):
    from openpyxl import Workbook

    original = tmp_path / "table.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+2"
    workbook.active["A2"] = "=WEBSERVICE(\"https://example.test\")"
    workbook.save(original)
    workbook.close()
    source = tmp_path / "cached.xlsx"
    with zipfile.ZipFile(original) as old, zipfile.ZipFile(source, "w") as new:
        for item in old.infolist():
            data = old.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v></v>", b"<v>3</v>", 1).replace(b"<v></v>", b"<v>9</v>", 1)
            new.writestr(item, data)
    output = tmp_path / "out.csv"
    converters.convert_file(source, output)
    assert output.read_text(encoding="utf-8").splitlines() == ["3", "9"]


def test_txt_docx_roundtrip_preserves_unicode_and_blank_lines(tmp_path):
    source = tmp_path / "notes.txt"
    text = "O‘zbekiston\n\nSardor va Madina\n"
    source.write_text(text, encoding="utf-8")
    document = tmp_path / "notes.docx"
    converters.convert_file(source, document)
    result = tmp_path / "result.txt"
    converters.convert_file(document, result)
    assert result.read_text(encoding="utf-8") == text


def test_docx_tables_keep_document_order(tmp_path):
    from docx import Document

    source = tmp_path / "input.docx"
    document = Document()
    document.add_paragraph("Avval")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Ism"
    table.cell(0, 1).text = "Sardor"
    document.add_paragraph("Keyin")
    document.save(source)
    result = tmp_path / "out.txt"
    converters.convert_file(source, result)
    assert result.read_text(encoding="utf-8").splitlines() == ["Avval", "Ism\tSardor", "Keyin"]


def test_pptx_text_and_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "O‘zbekcha taqdimot"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Jami"
    table.cell(0, 1).text = "100"
    source = tmp_path / "slides.pptx"
    presentation.save(source)
    result = tmp_path / "out.txt"
    converters.convert_file(source, result)
    text = result.read_text(encoding="utf-8")
    assert "Slayd 1" in text and "O‘zbekcha taqdimot" in text and "Jami\t100" in text


def test_scanned_pdf_does_not_claim_success(tmp_path):
    from pypdf import PdfWriter

    source = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(source)
    output = tmp_path / "text.txt"
    with pytest.raises(ValueError, match="OCR"):
        converters.convert_file(source, output)
    assert not output.exists()


def test_pdf_text_extraction(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    source = tmp_path / "text.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 100 Td (Jarvis test) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(source)
    output = tmp_path / "text.txt"
    converters.convert_file(source, output)
    assert "Jarvis test" in output.read_text(encoding="utf-8")


def test_zip_expansion_limit(tmp_path, monkeypatch):
    source = tmp_path / "bomb.docx"
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 1000)
    monkeypatch.setattr(converters, "MAX_EXPANDED_BYTES", 100)
    with pytest.raises(ValueError, match="hajm chegarasidan"):
        converters.convert_file(source, tmp_path / "out.txt")


def test_pdf_preflight_rejects_active_content(tmp_path):
    source = tmp_path / "evil.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("word/vbaProject.bin", b"not executable")
    with pytest.raises(ValueError, match="makros"):
        converters._check_office_active_content(source)


def test_pdf_preflight_rejects_external_links_but_allows_hyperlinks(tmp_path):
    source = tmp_path / "linked.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("word/_rels/document.xml.rels", '<Relationships><Relationship TargetMode="External" Type="image" Target="https://example.test/evil.png" /></Relationships>')
    with pytest.raises(ValueError, match="tashqi bog‘langan"):
        converters._check_office_active_content(source)
    safe = tmp_path / "safe.docx"
    with zipfile.ZipFile(safe, "w") as archive:
        archive.writestr("word/_rels/document.xml.rels", '<Relationships><Relationship TargetMode="External" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.test/" /></Relationships>')
    converters._check_office_active_content(safe)


@pytest.mark.parametrize("extension", [".docx", ".xlsx", ".pptx"])
def test_private_com_security_and_teardown(tmp_path, extension):
    source, destination = tmp_path / ("source" + extension), tmp_path / "result.pdf"
    application, document = MagicMock(), MagicMock()
    container = {".docx": application.Documents, ".xlsx": application.Workbooks, ".pptx": application.Presentations}[extension]
    def open_document(*args, **kwargs):
        assert application.AutomationSecurity == 3
        assert kwargs["ReadOnly"] is True
        if extension == ".xlsx":
            assert kwargs["UpdateLinks"] == 0
            assert application.EnableEvents is False
        return document
    container.Open.side_effect = open_document
    document.ExportAsFixedFormat.side_effect = lambda *args, **kwargs: destination.write_bytes(b"%PDF-1.7\n")
    dispatch = MagicMock(return_value=application)
    entered = []
    @contextmanager
    def owner(app):
        entered.append(app)
        yield
    assert "Microsoft" in _run_com(source, destination, dispatch, owner)
    assert entered == [application]
    application.Quit.assert_called_once()
    if extension == ".xlsx":
        application.CalculateFullRebuild.assert_called_once_with()
    if extension == ".pptx":
        document.Close.assert_called_once_with()
        assert document.Saved is True
    else:
        document.Close.assert_called_once_with(SaveChanges=False)


def test_private_com_teardown_on_export_error(tmp_path):
    application, document = MagicMock(), MagicMock()
    application.Documents.Open.return_value = document
    document.ExportAsFixedFormat.side_effect = RuntimeError("export error")
    @contextmanager
    def owner(app):
        yield
    with pytest.raises(RuntimeError, match="export error"):
        _run_com(tmp_path / "source.docx", tmp_path / "out.pdf", lambda _: application, owner)
    document.Close.assert_called_once_with(SaveChanges=False)
    application.Quit.assert_called_once()


def test_shared_office_instance_is_never_closed(tmp_path):
    application = MagicMock()
    @contextmanager
    def owner(app):
        raise ValueError("existing user process")
        yield
    with pytest.raises(ValueError, match="existing user process"):
        _run_com(tmp_path / "source.docx", tmp_path / "out.pdf", lambda _: application, owner)
    application.Quit.assert_not_called()
    application.Documents.Open.assert_not_called()


def test_pdf_worker_timeout_and_no_output(tmp_path, monkeypatch):
    from docx import Document

    source = tmp_path / "source.docx"
    Document().save(source)
    def timeout(*args, **kwargs):
        assert kwargs["timeout"] == converters.OFFICE_TIMEOUT_SECONDS
        assert "shell" not in kwargs
        raise subprocess.TimeoutExpired(args[0], 90)
    monkeypatch.setattr(converters.subprocess, "run", timeout)
    with pytest.raises(ValueError, match="vaqt chegarasidan"):
        converters.convert_file(source, tmp_path / "result.pdf")
    assert not (tmp_path / "result.pdf").exists()


def test_pdf_worker_success_copies_verified_output_exclusively(tmp_path, monkeypatch):
    from docx import Document

    source = tmp_path / "source.docx"
    Document().save(source)
    def success(command, **kwargs):
        Path(command[-1]).write_bytes(b"%PDF-1.7\nmock")
        return SimpleNamespace(returncode=0, stdout=json.dumps({"ok": True, "backend": "Word test"}))
    monkeypatch.setattr(converters.subprocess, "run", success)
    destination = tmp_path / "result.pdf"
    assert converters.convert_file(source, destination) == "Word test"
    assert destination.read_bytes().startswith(b"%PDF-")
    assert not list(tmp_path.glob("jarvis-pdf-*"))


def test_pdf_worker_error_is_clear(tmp_path, monkeypatch):
    from docx import Document

    source = tmp_path / "source.docx"
    Document().save(source)
    monkeypatch.setattr(converters.subprocess, "run", lambda *args, **kwargs: SimpleNamespace(returncode=1, stdout=json.dumps({"ok": False, "error": "Office aktivatsiyasini tekshiring"})))
    with pytest.raises(ValueError, match="aktivatsiyasini"):
        converters.convert_file(source, tmp_path / "result.pdf")


def test_worker_does_not_overwrite_late_created_output(tmp_path, monkeypatch):
    from docx import Document

    source = tmp_path / "source.docx"
    Document().save(source)
    destination = tmp_path / "result.pdf"
    def race(command, **kwargs):
        Path(command[-1]).write_bytes(b"%PDF-1.7\nmock")
        destination.write_bytes(b"another application created this")
        return SimpleNamespace(returncode=0, stdout=json.dumps({"ok": True}))
    monkeypatch.setattr(converters.subprocess, "run", race)
    with pytest.raises(FileExistsError):
        converters.convert_file(source, destination)
    assert destination.read_bytes() == b"another application created this"


def test_real_xlsx_rows_are_not_hidden_by_incorrect_dimension(tmp_path):
    from openpyxl import Workbook

    original, source = tmp_path / "original.xlsx", tmp_path / "dimension.xlsx"
    workbook = Workbook()
    workbook.active.append(["one"])
    workbook.active.append(["two"])
    workbook.save(original)
    workbook.close()
    with zipfile.ZipFile(original) as old, zipfile.ZipFile(source, "w") as new:
        for item in old.infolist():
            data = old.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:A2"', b'ref="A1"')
            new.writestr(item, data)
    destination = tmp_path / "out.csv"
    converters.convert_file(source, destination)
    assert destination.read_text(encoding="utf-8").splitlines() == ["one", "two"]


def test_repeated_shared_text_is_bounded(tmp_path, monkeypatch):
    from openpyxl import Workbook

    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    workbook.active.append(["long text", "long text", "long text"])
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(converters, "MAX_TEXT_CHARS", 10)
    with pytest.raises(ValueError, match="matni xavfsiz hajm"):
        converters.convert_file(source, tmp_path / "out.csv")
    assert not (tmp_path / "out.csv").exists()


def test_office_process_job_always_closes_own_handles(monkeypatch):
    import win32api
    import win32job

    application, job, process = MagicMock(), MagicMock(), MagicMock()
    monkeypatch.setattr(convert_worker, "_application_pid", lambda *args: 424242)
    monkeypatch.setattr(win32job, "CreateJobObject", lambda security, name: job)
    monkeypatch.setattr(win32job, "QueryInformationJobObject", lambda *args: {"BasicLimitInformation": {"LimitFlags": 0}})
    settings, assigned = MagicMock(), MagicMock()
    monkeypatch.setattr(win32job, "SetInformationJobObject", settings)
    monkeypatch.setattr(win32job, "AssignProcessToJobObject", assigned)
    monkeypatch.setattr(win32api, "OpenProcess", lambda *args: process)
    with pytest.raises(RuntimeError, match="failure"):
        with convert_worker._owned_process(application, {123}, ".docx"):
            raise RuntimeError("failure")
    assigned.assert_called_once_with(job, process)
    assert settings.call_args.args[2]["BasicLimitInformation"]["LimitFlags"] & win32job.JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
    job.Close.assert_called_once()
    process.Close.assert_called_once()


def test_office_job_rejects_existing_pid_without_quit(monkeypatch):
    monkeypatch.setattr(convert_worker, "_application_pid", lambda *args: 123)
    application = MagicMock()
    with pytest.raises(ValueError, match="mavjud jarayon"):
        with convert_worker._owned_process(application, {123}, ".docx"):
            pytest.fail("existing process ownership was accepted")
    application.Quit.assert_not_called()


def test_powerpoint_pid_must_be_new_and_unique(monkeypatch):
    monkeypatch.setattr(convert_worker, "_named_processes", lambda *args: {123, 456})
    with pytest.raises(ValueError, match="alohida yangi"):
        convert_worker._application_pid(MagicMock(), {123}, ".pptx")
    monkeypatch.setattr(convert_worker, "_named_processes", lambda *args: {456})
    assert convert_worker._application_pid(MagicMock(), {123}, ".pptx") == 456


def test_excel_pdf_preflight_accepts_local_sum_and_print_titles(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "sum.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["B2"], sheet["B3"], sheet["B4"] = 10, 20, "=SUM(B2:B3)"
    sheet.print_title_rows = "1:1"
    workbook.save(source)
    workbook.close()
    converters._check_office_active_content(source)


@pytest.mark.parametrize("formula", ['=HYPERLINK("https://example.test","click")', '=MyCustomFunction(A1)', '=INDIRECT("A1")', '=REGISTER.ID("dll","function")'])
def test_excel_pdf_preflight_rejects_unknown_or_dynamic_formulas(tmp_path, formula):
    from openpyxl import Workbook

    source = tmp_path / "unsafe.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 10
    workbook.active["A2"] = formula
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="Excel|formula"):
        converters._check_office_active_content(source)


def test_excel_pdf_preflight_checks_noncell_conditional_formulas(tmp_path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule

    source = tmp_path / "unsafe-conditional.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 10
    workbook.active.conditional_formatting.add("A1", FormulaRule(formula=["MyCustomFunction(A1)"]))
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="xavfsiz qayta hisoblash"):
        converters._check_office_active_content(source)


def test_excel_pdf_preflight_uses_correct_sheet_dimensions(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "sheets.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 1
    second = workbook.create_sheet("Long sheet")
    second["B10"] = 10
    second["B11"] = "=SUM(B1:B10)"
    workbook.save(source)
    workbook.close()
    converters._check_office_active_content(source)


def test_successful_export_survives_private_shutdown_proxy_error(tmp_path):
    application, document = MagicMock(), MagicMock()
    application.Workbooks.Open.return_value = document
    document.ExportAsFixedFormat.side_effect = lambda *args, **kwargs: (tmp_path / "out.pdf").write_bytes(b"%PDF-1.7\n")
    application.Workbooks.Add.return_value.Close.side_effect = RuntimeError("shutdown proxy unavailable")
    application.Quit.side_effect = RuntimeError("shutdown proxy unavailable")
    @contextmanager
    def owner(app):
        try:
            yield
        finally:
            app.job_kill_on_close()  # represents verified private job fallback
    assert "Microsoft Excel" in _run_com(tmp_path / "source.xlsx", tmp_path / "out.pdf", lambda _: application, owner)
    application.Quit.assert_called_once()
    application.job_kill_on_close.assert_called_once()


def test_com_retry_only_handles_transient_busy_errors(monkeypatch):
    import pythoncom

    monkeypatch.setattr(convert_worker.time, "sleep", lambda seconds: None)
    monkeypatch.setattr(pythoncom, "PumpWaitingMessages", lambda: None)
    operation = MagicMock(side_effect=[RuntimeError(-2147418111, "busy"), "result"])
    assert convert_worker._com_retry(operation) == "result"
    assert operation.call_count == 2
    operation = MagicMock(side_effect=RuntimeError("real export failure"))
    with pytest.raises(RuntimeError, match="real export failure"):
        convert_worker._com_retry(operation)
    operation.assert_called_once()
