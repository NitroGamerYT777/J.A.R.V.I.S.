"""Bounded, editable Office exports from data-only specifications.

No template, macro, script, executable expression or remote resource is loaded.
Validation and OOXML generation finish in memory before the output is created.
"""
from __future__ import annotations

from io import BytesIO
import math
from pathlib import Path
import re
from typing import Any


MAX_TEXT = 300_000
MAX_CELLS = 100_000
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_CELL = re.compile(r"\$?([A-Z]{1,3})\$?([1-9][0-9]{0,6})\Z", re.I)
_FUNCTIONS = frozenset("SUM AVERAGE MIN MAX COUNT COUNTA COUNTBLANK COUNTIF COUNTIFS SUMIF SUMIFS AVERAGEIF AVERAGEIFS IF IFERROR IFNA AND OR NOT ROUND ROUNDUP ROUNDDOWN ABS INT MOD POWER SQRT CONCAT CONCATENATE LEFT RIGHT MID LEN LOWER UPPER TRIM VALUE TEXT DATE YEAR MONTH DAY TODAY NOW VLOOKUP HLOOKUP INDEX MATCH XLOOKUP MEDIAN PRODUCT SUBTOTAL CEILING FLOOR LARGE SMALL RANK SUMPRODUCT".split())


def _obj(value: Any, label: str, allowed: set[str]) -> dict:
    if not isinstance(value, dict) or any(not isinstance(k, str) for k in value):
        raise ValueError(f"{label}: JSON obyekt bo‘lishi kerak.")
    unknown = set(value) - allowed
    if unknown:
        raise ValueError(f"{label}: noma’lum maydon: {', '.join(sorted(unknown))}.")
    return value


def _text(value: Any, label: str, budget: list[int], maximum: int = 6000, *, required: bool = False) -> str:
    if not isinstance(value, str) or _CONTROL.search(value):
        raise ValueError(f"{label}: oddiy matn bo‘lishi kerak.")
    if len(value) > maximum or (required and not value.strip()):
        raise ValueError(f"{label}: matn bo‘sh yoki {maximum} belgidan uzun.")
    budget[0] += len(value)
    if budget[0] > MAX_TEXT:
        raise ValueError(f"Hujjat matni {MAX_TEXT} belgidan oshmasin.")
    return value


def _items(value: Any, label: str, maximum: int, minimum: int = 0) -> list:
    if not isinstance(value, list) or not minimum <= len(value) <= maximum:
        raise ValueError(f"{label}: {minimum}–{maximum} elementli ro‘yxat bo‘lishi kerak.")
    return value


def _scalar(value: Any, label: str, budget: list[int], maximum: int = 4000) -> None:
    if isinstance(value, str):
        _text(value, label, budget, maximum)
    elif value is not None and type(value) not in (bool, int, float):
        raise ValueError(f"{label}: faqat matn, son, mantiqiy qiymat yoki null mumkin.")
    elif type(value) in (int, float) and (abs(value) > 1e100 or not math.isfinite(value)):
        raise ValueError(f"{label}: chekli va odatiy kattalikdagi son kerak.")


def _table(value: Any, label: str, budget: list[int], rows: int, cols: int, *, cell_text: int = 1000) -> dict:
    table = _obj(value, label, {"headers", "rows"})
    headers = _items(table.get("headers"), f"{label} sarlavhalari", cols, 1)
    for header in headers:
        _text(header, f"{label} ustun nomi", budget, min(cell_text, 200), required=True)
    for row in _items(table.get("rows", []), f"{label} qatorlari", rows):
        if not isinstance(row, list) or len(row) != len(headers):
            raise ValueError(f"{label}: har qator {len(headers)} katakdan iborat bo‘lishi kerak.")
        for item in row:
            _scalar(item, f"{label} katagi", budget, cell_text)
    return table


def _validate_docx(spec: dict, budget: list[int]) -> None:
    _obj(spec, "Word hujjati", {"title", "subtitle", "sections"})
    _text(spec.get("title"), "Hujjat nomi", budget, 220, required=True)
    _text(spec.get("subtitle", ""), "Taglavha", budget, 500)
    cells = 0
    for section in _items(spec.get("sections"), "Bo‘limlar", 80, 1):
        _obj(section, "Bo‘lim", {"heading", "paragraphs", "bullets", "table"})
        _text(section.get("heading", ""), "Bo‘lim nomi", budget, 220)
        for paragraph in _items(section.get("paragraphs", []), "Abzatslar", 150):
            _text(paragraph, "Abzats", budget)
        for bullet in _items(section.get("bullets", []), "Ro‘yxat", 100):
            _text(bullet, "Ro‘yxat matni", budget, 1200)
        if "table" in section:
            _table(section["table"], "Word jadvali", budget, 300, 12)
            cells += len(section["table"]["headers"]) * (len(section["table"].get("rows", [])) + 1)
            if cells > 20_000:
                raise ValueError("Word hujjatidagi jadvallar jami 20000 katakdan oshmasin.")
        if not any(section.get(k) for k in ("heading", "paragraphs", "bullets", "table")):
            raise ValueError("Word bo‘limi bo‘sh bo‘lmasin.")


def _cell_address(address: str, rows: int, cols: int) -> None:
    from openpyxl.utils.cell import column_index_from_string

    match = _CELL.fullmatch(address)
    if not match or int(match[2]) > rows or column_index_from_string(match[1]) > cols:
        raise ValueError("Formula manzili jadval chegaralari ichida bo‘lishi kerak.")


def _formula(formula: Any, sheets: dict[str, tuple[int, int]], current: str) -> None:
    from openpyxl.formula.tokenizer import Tokenizer, TokenizerError

    if not isinstance(formula, str) or not formula.startswith("=") or len(formula) > 1500:
        raise ValueError("Formula '=' bilan boshlanuvchi, 1500 belgigacha bo‘lgan matn bo‘lsin.")
    if any(c in formula for c in "[]|{};\r\n") or _CONTROL.search(formula):
        raise ValueError("Formulada tashqi havola, DDE yoki massiv buyruqlari mumkin emas.")
    try:
        tokens = Tokenizer(formula).items
    except (TokenizerError, IndexError, ValueError) as exc:
        raise ValueError("Excel formulasining tuzilishi noto‘g‘ri.") from exc
    if not tokens:
        raise ValueError("Excel formulasi bo‘sh bo‘lmasin.")
    stack: list[str] = []
    expected_operand = True
    previous = None
    for token in tokens:
        if token.type == "WHITE-SPACE":
            continue
        if token.type == "FUNC" and token.subtype == "OPEN":
            if not expected_operand:
                raise ValueError("Formula ifodalari orasida amal belgisi yetishmaydi.")
            name = token.value[:-1].upper()
            if name not in _FUNCTIONS:
                raise ValueError(f"{name}: ushbu Excel funksiyasiga xavfsizlik uchun ruxsat berilmagan.")
            stack.append("FUNC")
        elif token.type == "FUNC" and token.subtype == "CLOSE":
            if not stack or stack.pop() != "FUNC" or (expected_operand and not (previous and previous.type == "FUNC" and previous.subtype == "OPEN")):
                raise ValueError("Formula qavslari yoki argumentlari noto‘g‘ri.")
            expected_operand = False
        elif token.type == "PAREN" and token.subtype == "OPEN":
            if not expected_operand:
                raise ValueError("Formula ifodalari orasida amal belgisi yetishmaydi.")
            stack.append("PAREN")
        elif token.type == "PAREN" and token.subtype == "CLOSE":
            if not stack or stack.pop() != "PAREN" or expected_operand:
                raise ValueError("Formula qavslari mos emas.")
            expected_operand = False
        elif token.type == "OPERAND":
            if not expected_operand:
                raise ValueError("Formula ifodalari orasida amal belgisi yetishmaydi.")
            if token.subtype == "RANGE":
                reference = token.value
                sheet_name = current
                if "!" in reference:
                    sheet_name, reference = reference.rsplit("!", 1)
                    if sheet_name.startswith("'") and sheet_name.endswith("'"):
                        sheet_name = sheet_name[1:-1].replace("''", "'")
                    if sheet_name.casefold() not in sheets:
                        raise ValueError("Formula faqat shu fayldagi mavjud varaqlarga murojaat qilsin.")
                bounds = sheets[sheet_name.casefold()]
                parts = reference.split(":")
                if len(parts) > 2:
                    raise ValueError("Formula kataklar diapazoni noto‘g‘ri.")
                for address in parts:
                    _cell_address(address, *bounds)
            elif token.subtype not in {"NUMBER", "TEXT", "LOGICAL", "ERROR"}:
                raise ValueError("Formulada noma’lum ifoda bor.")
            expected_operand = False
        elif token.type == "OPERATOR-PREFIX":
            if not expected_operand or token.value not in {"+", "-"}:
                raise ValueError("Formula amal belgisi noto‘g‘ri.")
        elif token.type == "OPERATOR-INFIX":
            if expected_operand or token.value not in {"+", "-", "*", "/", "^", "&", "=", "<", ">", "<=", ">=", "<>"}:
                raise ValueError("Formula amal belgisi noto‘g‘ri.")
            expected_operand = True
        elif token.type == "OPERATOR-POSTFIX":
            if expected_operand or token.value != "%":
                raise ValueError("Formula foiz ifodasi noto‘g‘ri.")
        elif token.type == "SEP":
            if expected_operand or not stack or stack[-1] != "FUNC" or token.subtype != "ARG":
                raise ValueError("Formula argumentlari noto‘g‘ri.")
            expected_operand = True
        else:
            raise ValueError("Formula turi qo‘llab-quvvatlanmaydi.")
        previous = token
    if stack or expected_operand:
        raise ValueError("Formula tugallanmagan yoki qavslari mos emas.")


def _validate_xlsx(spec: dict, budget: list[int]) -> None:
    _obj(spec, "Excel hujjati", {"title", "sheets"})
    _text(spec.get("title"), "Hujjat nomi", budget, 220, required=True)
    sheets = _items(spec.get("sheets"), "Varaqlar", 12, 1)
    names: dict[str, tuple[int, int]] = {}
    cells = 0
    for sheet in sheets:
        _obj(sheet, "Varaq", {"name", "headers", "rows", "formulas", "chart"})
        name = _text(sheet.get("name"), "Varaq nomi", budget, 31, required=True)
        if re.search(r"[\[\]:*?/\\]", name) or name.startswith("'") or name.endswith("'") or name.casefold() in names:
            raise ValueError("Excel varaq nomlari takrorlanmasin va []:*?/\\ belgilarini saqlamasin.")
        _table({"headers": sheet.get("headers"), "rows": sheet.get("rows", [])}, "Excel jadvali", budget, 5000, 100, cell_text=4000)
        headers = sheet["headers"]
        if len({header.casefold() for header in headers}) != len(headers):
            raise ValueError("Excel ustun nomlari takrorlanmasin.")
        rows = len(sheet.get("rows", [])) + 1
        names[name.casefold()] = (rows, len(headers))
        cells += rows * len(headers)
        if cells > MAX_CELLS:
            raise ValueError(f"Excel fayli jami {MAX_CELLS} katakdan oshmasin.")
    for sheet in sheets:
        bounds = names[sheet["name"].casefold()]
        addresses: set[str] = set()
        for entry in _items(sheet.get("formulas", []), "Formulalar", 5000):
            _obj(entry, "Formula", {"cell", "formula"})
            address = _text(entry.get("cell"), "Formula katagi", budget, 12, required=True).upper()
            _cell_address(address, *bounds)
            if address.replace("$", "") in addresses or re.search(r"\$?1$", address):
                raise ValueError("Formula katagi takrorlangan yoki sarlavha qatorida joylashgan.")
            addresses.add(address.replace("$", ""))
            formula = _text(entry.get("formula"), "Formula matni", budget, 1500, required=True)
            _formula(formula, names, sheet["name"])
        if "chart" in sheet:
            chart = _obj(sheet["chart"], "Diagramma", {"type", "title", "category_column", "value_columns"})
            if not isinstance(chart.get("type"), str) or chart["type"] not in {"bar", "line", "pie"} or bounds[0] < 2:
                raise ValueError("Diagramma turi bar, line yoki pie bo‘lsin va ma’lumot qatori mavjud bo‘lsin.")
            _text(chart.get("title", ""), "Diagramma nomi", budget, 200)
            category = chart.get("category_column", 1)
            values = _items(chart.get("value_columns", [2]), "Diagramma qiymat ustunlari", 10, 1)
            if type(category) is not int or not 1 <= category <= bounds[1]:
                raise ValueError("Diagramma kategoriya ustuni mavjud bo‘lishi kerak.")
            if any(type(col) is not int or not 1 <= col <= bounds[1] or col == category for col in values) or len(set(values)) != len(values):
                raise ValueError("Diagramma qiymat ustunlari mavjud va takrorlanmagan bo‘lsin.")
            if chart["type"] == "pie" and len(values) != 1:
                raise ValueError("Doiraviy diagramma bitta qiymat ustunini qabul qiladi.")


def _lines(text: str, chars: int) -> int:
    return sum(max(1, math.ceil(len(line) / chars)) for line in text.split("\n"))


def _slide_table_heights(slide: dict) -> list[float]:
    table = slide["table"]
    cols = len(table["headers"])
    width = 11.7 / cols
    chars = max(12, int(width * 8))
    rows = [table["headers"], *table.get("rows", [])]
    heights = [max(0.46, max(_lines(_display(cell), chars) for cell in row) * 0.22 + 0.16) for row in rows]
    allowed = 3.75 if slide.get("body") or slide.get("bullets") else 4.75
    if sum(heights) > allowed:
        raise ValueError("Slayd jadvali sig‘maydi: qatorlarni yoki kataklardagi matnni qisqartiring.")
    return heights


def _validate_pptx(spec: dict, budget: list[int]) -> None:
    _obj(spec, "PowerPoint hujjati", {"title", "slides"})
    _text(spec.get("title"), "Hujjat nomi", budget, 220, required=True)
    for slide in _items(spec.get("slides"), "Slaydlar", 40, 1):
        _obj(slide, "Slayd", {"title", "body", "bullets", "notes", "table"})
        _text(slide.get("title"), "Slayd nomi", budget, 100, required=True)
        body = _text(slide.get("body", ""), "Slayd matni", budget, 1100)
        bullets = _items(slide.get("bullets", []), "Slayd ro‘yxati", 8)
        for bullet in bullets:
            _text(bullet, "Slayd ro‘yxat matni", budget, 240)
        _text(slide.get("notes", ""), "Taqdimot izohi", budget, 6000)
        line_count = (_lines(body, 82) if body else 0) + sum(_lines(b, 78) + 0.4 for b in bullets)
        if line_count > (4 if "table" in slide else 14):
            raise ValueError("Slayd matni sig‘maydi: matnni qisqartiring yoki bir nechta slaydga ajrating.")
        if "table" in slide:
            _table(slide["table"], "Slayd jadvali", budget, 8, 6, cell_text=100)
            _slide_table_heights(slide)


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Ha" if value else "Yo‘q"
    return str(value)


def _docx(spec: dict, stream: BytesIO) -> None:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    doc.core_properties.title = spec["title"]
    doc.core_properties.author = "JARVIS"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.1)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor.from_string("000000")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.size = Pt(28)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.add_paragraph(spec["title"], "Title")
    if spec.get("subtitle"):
        doc.add_paragraph(spec["subtitle"], "Subtitle")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run("JARVIS  ·  ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("667085")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for content in spec["sections"]:
        if content.get("heading"):
            doc.add_paragraph(content["heading"], "Heading 1")
        for paragraph in content.get("paragraphs", []):
            doc.add_paragraph(paragraph)
        for bullet in content.get("bullets", []):
            doc.add_paragraph(bullet, "List Bullet")
        if "table" not in content:
            continue
        data = content["table"]
        table = doc.add_table(rows=1, cols=len(data["headers"]))
        table.autofit = False
        table.style = "Table Grid"
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            for key, value in {"val": "single", "sz": "5", "color": "D9D9D9"}.items():
                border.set(qn(f"w:{key}"), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        rows = [data["headers"], *data.get("rows", [])]
        for index, row_data in enumerate(rows):
            cells = table.rows[0].cells if index == 0 else table.add_row().cells
            for cell, value in zip(cells, row_data):
                cell.width = Cm(16.8 / len(cells))
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = _display(value)
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "203B5F" if index == 0 else ("EFF4F9" if index % 2 == 0 else "FFFFFF"))
                props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement(f"w:{side}")
                    margin.set(qn("w:w"), "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        if index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            no_split = OxmlElement("w:cantSplit")
            table.rows[index]._tr.get_or_add_trPr().append(no_split)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    doc.save(stream)


def _xlsx(spec: dict, stream: BytesIO) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties

    book = Workbook()
    book.remove(book.active)
    book.properties.title = spec["title"]
    book.properties.creator = "JARVIS"
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    for index, data in enumerate(spec["sheets"], 1):
        sheet = book.create_sheet(data["name"])
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "147D92"
        sheet.freeze_panes = "A2"
        for row_index, row in enumerate([data["headers"], *data.get("rows", [])], 1):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row_index, column, value)
                if isinstance(value, str):
                    cell.data_type = "s"  # AI/user strings never become executable formulas.
                cell.font = Font(name="Calibri", size=11, color="FFFFFF" if row_index == 1 else "172B4D", bold=row_index == 1)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row_index == 1:
                    cell.fill = PatternFill("solid", fgColor="203B5F")
                elif type(value) in (float, int):
                    cell.number_format = "#,##0.00" if isinstance(value, float) else "#,##0"
            sheet.row_dimensions[row_index].height = 30 if row_index == 1 else 24
        for formula in data.get("formulas", []):
            sheet[formula["cell"].replace("$", "").upper()] = formula["formula"]
        end = f"{get_column_letter(len(data['headers']))}{sheet.max_row}"
        sheet.auto_filter.ref = f"A1:{end}"
        if sheet.max_row > 1:
            table = Table(displayName=f"JARVISTable{index}", ref=f"A1:{end}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        for column in range(1, len(data["headers"]) + 1):
            values = [len(_display(sheet.cell(row, column).value)) for row in range(1, min(sheet.max_row, 101) + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(45, max(14, max(values, default=10) + 3))
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        if "chart" in data:
            options = data["chart"]
            chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[options["type"]]()
            chart.title = options.get("title") or spec["title"]
            chart.style = 10
            for column in options.get("value_columns", [2]):
                chart.add_data(Reference(sheet, min_col=column, min_row=1, max_row=sheet.max_row), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=options.get("category_column", 1), min_row=2, max_row=sheet.max_row))
            chart.width, chart.height = 17, 10
            sheet.add_chart(chart, f"{get_column_letter(len(data['headers']) + 2)}2")
    book.save(stream)


def _pptx(spec: dict, stream: BytesIO) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = spec["title"]
    deck.core_properties.author = "JARVIS"

    def textbox(slide, text, left, top, width, height, size, color="172B4D", bold=False):
        frame = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.text = text
        for paragraph in frame.paragraphs:
            paragraph.font.name = "Calibri"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
        return frame

    for index, content in enumerate(spec["slides"], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        textbox(slide, content["title"], .8, .55, 11.7, 1.05, 32, "203B5F", True)
        body, bullets = content.get("body", ""), content.get("bullets", [])
        if body or bullets:
            frame = textbox(slide, body, .85, 1.8, 11.6, 1.1 if "table" in content else 4.7, 20)
            if not body:
                frame.clear()
            for bullet_index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if not body and bullet_index == 0 else frame.add_paragraph()
                paragraph.text = bullet
                paragraph.font.name = "Calibri"
                paragraph.font.size = Pt(20)
                paragraph.font.color.rgb = RGBColor.from_string("172B4D")
                paragraph.space_after = Pt(8)
                paragraph.level = 0
                bullet_marker = OxmlElement("a:buChar")
                bullet_marker.set("char", "•")
                paragraph._p.get_or_add_pPr().append(bullet_marker)
        if "table" in content:
            data = content["table"]
            heights = _slide_table_heights(content)
            top = 3.0 if body or bullets else 1.85
            table = slide.shapes.add_table(len(heights), len(data["headers"]), Inches(.8), Inches(top), Inches(11.7), Inches(sum(heights))).table
            rows = [data["headers"], *data.get("rows", [])]
            for row_index, row in enumerate(rows):
                table.rows[row_index].height = Inches(heights[row_index])
                for column, value in enumerate(row):
                    cell = table.cell(row_index, column)
                    cell.text = _display(value)
                    cell.margin_left = cell.margin_right = Inches(.1)
                    cell.margin_top = cell.margin_bottom = Inches(.06)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string("203B5F" if row_index == 0 else ("EDF3F8" if row_index % 2 == 0 else "FFFFFF"))
                    cell.text_frame.word_wrap = True
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = "Calibri"
                        paragraph.font.size = Pt(14)
                        paragraph.font.bold = row_index == 0
                        paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else "172B4D")
        textbox(slide, "JARVIS", .8, 7.05, 2, .2, 9, "667085")
        textbox(slide, f"{index:02d} / {len(spec['slides']):02d}", 11.5, 7.05, 1, .2, 9, "667085")
        if content.get("notes"):
            slide.notes_slide.notes_text_frame.text = content["notes"]
    deck.save(stream)


def _validate_txt(spec, budget):
    _obj(spec, 'Matn fayli', {'title', 'text'})
    _text(spec.get('title'), 'Hujjat nomi', budget, 220, required=True)
    _text(spec.get('text'), 'Hujjat matni', budget, 100000, required=True)


def _txt(spec, stream):
    stream.write(spec['text'].encode('utf-8'))


def validate_document(spec: dict, suffix: str) -> None:
    validators = {".docx": _validate_docx, ".xlsx": _validate_xlsx, ".pptx": _validate_pptx, '.txt': _validate_txt}
    if suffix not in validators:
        raise ValueError("Faqat .docx, .xlsx, .pptx va .txt fayllarini yaratish mumkin.")
    validators[suffix](spec, [0])


def render_document(spec: dict, destination: Path) -> None:
    """Create a new DOCX/XLSX/PPTX, without ever overwriting an existing file."""
    destination = Path(destination)
    suffix = destination.suffix.lower()
    renderers = {".docx": (_validate_docx, _docx), ".xlsx": (_validate_xlsx, _xlsx), ".pptx": (_validate_pptx, _pptx), '.txt': (_validate_txt, _txt)}
    if suffix not in renderers:
        raise ValueError("Faqat .docx, .xlsx, .pptx va .txt fayllarini yaratish mumkin.")
    if destination.exists():
        raise ValueError("Bunday fayl mavjud. Boshqa nom tanlang; mavjud fayl almashtirilmaydi.")
    if not destination.parent.is_dir():
        raise ValueError("Saqlash papkasi mavjud emas.")
    validate, render = renderers[suffix]
    validate(spec, [0])
    stream = BytesIO()
    render(spec, stream)
    try:
        output = destination.open("xb")
    except FileExistsError as exc:
        raise ValueError("Bunday fayl mavjud; mavjud fayl almashtirilmaydi.") from exc
    try:
        with output:
            output.write(stream.getbuffer())
    except BaseException:
        destination.unlink(missing_ok=True)
        raise
