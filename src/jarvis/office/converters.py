"""Bounded, non-overwriting Office conversions; importing this module runs no Office code."""

from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile


MAX_SOURCE_BYTES = 30 * 1024 * 1024
MAX_EXPANDED_BYTES = 120 * 1024 * 1024
MAX_CELLS = 500_000
MAX_ROWS = 100_000
MAX_TEXT_CHARS = 20_000_000
OFFICE_TIMEOUT_SECONDS = 90
_CONVERSIONS = {
    ".csv": (".xlsx",),
    ".xlsx": (".csv", ".pdf"),
    ".txt": (".docx",),
    ".docx": (".txt", ".pdf"),
    ".pptx": (".txt", ".pdf", ".docx"),
    ".pdf": (".txt",),
}


def supported_conversions() -> dict[str, tuple[str, ...]]:
    return dict(_CONVERSIONS)


def _check_source(source: Path) -> None:
    if not source.is_file():
        raise ValueError("Manba fayl topilmadi.")
    if source.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("Manba fayl 30 MB chegaradan katta.")
    if source.suffix.lower() in {".docx", ".xlsx", ".pptx"}:
        try:
            with zipfile.ZipFile(source) as archive:
                entries = archive.infolist()
                if len(entries) > 10_000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                    raise ValueError("Office fayli ochilganda xavfsiz hajm chegarasidan oshadi.")
                if any(item.flag_bits & 1 for item in entries):
                    raise ValueError("Parollangan Office arxivi qo‘llab-quvvatlanmaydi.")
        except zipfile.BadZipFile as exc:
            raise ValueError("Office fayli buzilgan yoki formati noto‘g‘ri.") from exc


def _safe_cell(value: object) -> object:
    """Do not let a CSV/string cell turn into an Excel formula when opened."""
    if isinstance(value, str) and (
        value.startswith(("\t", "\r", "\n")) or value.lstrip().startswith(("=", "+", "-", "@"))
    ):
        return "'" + value
    return value


def _read_text(source: Path) -> str:
    try:
        return source.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Matn fayli UTF-8 kodlashda bo‘lishi kerak.") from exc


def _write_text(destination: Path, content: str) -> None:
    with destination.open("x", encoding="utf-8", newline="") as stream:
        stream.write(content)


def _csv_to_xlsx(source: Path, destination: Path) -> str:
    from openpyxl import Workbook

    text = _read_text(source)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ma’lumotlar"
    count = 0
    try:
        for number, row in enumerate(csv.reader(io.StringIO(text, newline=""), dialect), start=1):
            count += len(row)
            if number > MAX_ROWS or count > MAX_CELLS or len(row) > 16_384:
                raise ValueError("Jadval xavfsiz qator yoki katak soni chegarasidan oshdi.")
            if any(len(value) > 32_767 for value in row):
                raise ValueError("Excel katagi 32 767 belgidan oshmasligi kerak; ma’lumot qisqartirilmadi.")
            worksheet.append([_safe_cell(value) for value in row])
        with destination.open("xb") as stream:
            workbook.save(stream)
    finally:
        workbook.close()
    return "openpyxl (CSV → XLSX; formulaga o‘xshash matnlar xavfsizlashtirildi)"


def _xlsx_to_csv(source: Path, destination: Path, sheet: str | None) -> str:
    from openpyxl import load_workbook

    formulas = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    values = None
    try:
        if sheet is None and len(formulas.sheetnames) != 1:
            raise ValueError("CSV uchun bitta varaqni tanlang: " + ", ".join(formulas.sheetnames))
        name = sheet if sheet is not None else formulas.sheetnames[0]
        if name not in formulas.sheetnames:
            raise ValueError("Bunday Excel varag‘i yo‘q. Mavjudlari: " + ", ".join(formulas.sheetnames))
        worksheet = formulas[name]
        if (worksheet.max_row or 0) > MAX_ROWS or (worksheet.max_row or 0) * (worksheet.max_column or 0) > MAX_CELLS:
            raise ValueError("Excel varag‘i xavfsiz katak soni chegarasidan oshdi.")
        values = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        # A forged/incorrect <dimension> must not silently hide real rows.
        worksheet.reset_dimensions()
        values[name].reset_dimensions()
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        cell_count = text_count = 0
        for row_number, (formula_row, value_row) in enumerate(zip(worksheet.iter_rows(), values[name].iter_rows()), start=1):
            cell_count += len(formula_row)
            if row_number > MAX_ROWS or cell_count > MAX_CELLS:
                raise ValueError("Excel varag‘i xavfsiz qator yoki katak soni chegarasidan oshdi.")
            row = []
            for cell, cached in zip(formula_row, value_row):
                value = cell.value
                if cell.data_type == "f":
                    if cached.value is None:
                        raise ValueError(
                            f"{name}!{cell.coordinate}: formula hisoblangan qiymati yo‘q. "
                            "Faylni Excel’da hisoblab saqlang, keyin qayta aylantiring."
                        )
                    value = cached.value
                row.append(_safe_cell(value))
                text_count += len(str(value)) if value is not None else 0
                if text_count > MAX_TEXT_CHARS:
                    raise ValueError("Excel matni xavfsiz hajm chegarasidan oshdi; natija qisqartirilmadi.")
            writer.writerow(row)
        _write_text(destination, output.getvalue())
    finally:
        formulas.close()
        if values is not None:
            values.close()
    return "openpyxl (XLSX → CSV; hisoblangan qiymatlar, formulalar bajarilmadi)"


def _docx_to_text(source: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(source)
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            parts.append(Paragraph(child, document).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _pptx_to_text(source: Path) -> str:
    from pptx import Presentation

    def shape_text(shape: object) -> list[str]:
        if hasattr(shape, "shapes"):
            return [text for child in shape.shapes for text in shape_text(child)]
        if shape.has_text_frame:
            return [shape.text_frame.text]
        if shape.has_table:
            return ["\t".join(cell.text for cell in row.cells) for row in shape.table.rows]
        return []

    presentation = Presentation(source)
    return "\n\n".join(
        f"Slayd {number}\n" + "\n".join(text for shape in slide.shapes for text in shape_text(shape))
        for number, slide in enumerate(presentation.slides, start=1)
    )


def _pptx_to_docx(source: Path, destination: Path) -> str:
    """Editable slide text and tables, not a claim to reproduce slide artwork."""
    from docx import Document
    from docx.shared import Pt
    from pptx import Presentation

    presentation = Presentation(source)
    if not 1 <= len(presentation.slides) <= 500:
        raise ValueError("Wordga konvertatsiya uchun 1–500 ta slayd kerak.")
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.core_properties.title = presentation.core_properties.title or source.stem
    document.add_heading(document.core_properties.title, 0)
    text_count = cell_count = 0
    meaningful = False

    def shapes(items):
        for shape in items:
            if hasattr(shape, "shapes"):
                yield from shapes(shape.shapes)
            else:
                yield shape

    for number, slide in enumerate(presentation.slides, 1):
        if number > 1:
            document.add_page_break()
        document.add_heading(f"Slayd {number}", 1)
        for shape in shapes(slide.shapes):
            if shape.has_table:
                source_table = shape.table
                rows, columns = len(source_table.rows), len(source_table.columns)
                cell_count += rows*columns
                if rows > 1000 or columns > 50 or cell_count > 50000:
                    raise ValueError("Slayd jadvallari xavfsiz Word hajmi chegarasidan oshdi.")
                table = document.add_table(rows=rows, cols=columns)
                table.style = "Light Shading Accent 1"
                for row_index, row in enumerate(source_table.rows):
                    for column, cell in enumerate(row.cells):
                        text = cell.text.replace("\v", "\n")
                        table.cell(row_index, column).text = text
                        text_count += len(text)
                        meaningful |= bool(text.strip())
            elif shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.replace("\v", "\n")
                    if text.strip():
                        document.add_paragraph(text)
                        text_count += len(text)
                        meaningful = True
            if text_count > 2_000_000:
                raise ValueError("Slayd matni Word konvertatsiyasi chegarasidan oshdi.")
    if not meaningful:
        raise ValueError("Slaydlarda olinadigan matn/jadval yo'q. Faqat rasmli slaydlar uchun OCR kerak.")
    stream = io.BytesIO()
    document.save(stream)
    with destination.open("xb") as output:
        output.write(stream.getbuffer())
    return "python-pptx + python-docx (PPTX → DOCX; slayd matni va jadvallar; rasmlar, izohlar va slayd bezagi ko'chirilmaydi)"


def _pdf_to_text(source: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(source)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF parol bilan himoyalangan.")
    if len(reader.pages) > 1000:
        raise ValueError("PDF 1 000 sahifadan katta.")
    parts = []
    total = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        total += len(text)
        if total > MAX_TEXT_CHARS:
            raise ValueError("PDF matni xavfsiz hajm chegarasidan oshdi.")
        parts.append(text)
    content = "\n\n".join(parts)
    if not content.strip():
        raise ValueError("PDF’da ajratib olinadigan matn yo‘q. Skanerlangan PDF uchun OCR kerak; hozir OCR bajarilmaydi.")
    return content


def _check_office_active_content(source: Path) -> None:
    """Reject active/linked content before a desktop Office application sees it."""
    with zipfile.ZipFile(source) as archive:
        for entry in archive.infolist():
            name = entry.filename.lower()
            if any(part in name for part in (
                "vbaproject", "/embeddings/", "/activex/", "/macrosheets/", "/externallinks/", "/querytables/",
            )) or name == "xl/connections.xml":
                raise ValueError("PDF konvertatsiyasi: makros, ichki dastur yoki tashqi ma’lumot manbali fayl rad etildi.")
            if name.endswith(".rels"):
                root = ET.fromstring(archive.read(entry))
                for relation in root:
                    if relation.attrib.get("TargetMode", "").lower() == "external" and not relation.attrib.get("Type", "").endswith("/hyperlink"):
                        raise ValueError("PDF konvertatsiyasi: tashqi bog‘langan kontentni avval fayl ichiga joylang.")
            if name.startswith("word/") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(entry))
                instructions = " ".join(
                    (node.text or "") for node in root.iter() if node.tag.endswith("}instrText")
                )
                if re.search(r"\b(DDEAUTO|DDE|INCLUDETEXT|INCLUDEPICTURE|DATABASE|LINK)\b", instructions, re.I):
                    raise ValueError("PDF konvertatsiyasi: tashqi yoki DDE Word maydonlari xavfsizlik uchun rad etildi.")
            if name.startswith("xl/") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(entry))
                for node in root.iter():
                    if node.tag.endswith("}f") or node.tag.endswith("}definedName"):
                        formula = node.text or ""
                        if re.search(r"\[|\||https?://|\b(WEBSERVICE|RTD|DDE|CALL|REGISTER|EXEC)\s*\(", formula, re.I):
                            raise ValueError("PDF konvertatsiyasi: tashqi yoki bajariladigan Excel formulasi rad etildi.")
    if source.suffix.lower() == ".xlsx":
        _check_xlsx_recalculation(source)


def _check_xlsx_recalculation(source: Path) -> None:
    """Only known local formulas may be recalculated by the private Excel host."""
    from openpyxl import load_workbook
    from openpyxl.formula.tokenizer import Tokenizer
    from openpyxl.utils.cell import column_index_from_string
    from .renderers import _formula

    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        bounds = {sheet.title.casefold(): (sheet.max_row or 1, sheet.max_column or 1) for sheet in workbook.worksheets}
        sheet_paths = {sheet._worksheet_path: sheet.title for sheet in workbook.worksheets}
        sheet_names = workbook.sheetnames
        if not bounds or sum(rows * cols for rows, cols in bounds.values()) > MAX_CELLS:
            raise ValueError("Excel hisoblash kataklari xavfsiz chegaradan oshdi.")
        current = workbook.sheetnames[0]
    finally:
        workbook.close()

    def print_titles(expression: str) -> None:
        for token in Tokenizer("=" + expression).items:
            if token.type == "SEP" and token.value == ",":
                continue
            if token.type != "OPERAND" or token.subtype != "RANGE" or "!" not in token.value:
                raise ValueError("Excel chop etish qator/ustun nomlari oddiy mahalliy diapazon bo‘lsin.")
            sheet_name, reference = token.value.rsplit("!", 1)
            sheet_name = sheet_name.strip("'").replace("''", "'").casefold()
            if sheet_name not in bounds:
                raise ValueError("Excel chop etish varag‘i mavjud emas.")
            rows, columns = bounds[sheet_name]
            if re.fullmatch(r"\$?[1-9][0-9]*:\$?[1-9][0-9]*", reference):
                if any(int(value.replace("$", "")) > rows for value in reference.split(":")):
                    raise ValueError("Excel chop etish qatorlari jadval chegarasidan tashqarida.")
            elif re.fullmatch(r"\$?[A-Z]{1,3}:\$?[A-Z]{1,3}", reference, re.I):
                if any(column_index_from_string(value.replace("$", "")) > columns for value in reference.split(":")):
                    raise ValueError("Excel chop etish ustunlari jadval chegarasidan tashqarida.")
            else:
                raise ValueError("Excel chop etish diapazoni xavfsiz statik qator/ustun bo‘lsin.")

    with zipfile.ZipFile(source) as archive:
        for entry in archive.infolist():
            if not entry.filename.startswith("xl/") or not entry.filename.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(entry))
            for node in root.iter():
                tag = node.tag.rsplit("}", 1)[-1]
                formula = None
                if tag in {"f", "formula", "formula1", "formula2", "calculatedColumnFormula", "totalsRowFormula", "definedName"}:
                    formula = node.text
                elif tag == "cfvo" and node.attrib.get("type") == "formula":
                    formula = node.attrib.get("val")
                if not formula:
                    continue
                try:
                    if tag == "definedName" and node.attrib.get("name") == "_xlnm.Print_Titles":
                        print_titles(formula)
                    else:
                        formula_sheet = sheet_paths.get(entry.filename, current)
                        if tag == "definedName" and node.attrib.get("localSheetId", "").isdigit():
                            index = int(node.attrib["localSheetId"])
                            if index >= len(sheet_names):
                                raise ValueError("Excel formula varag‘i mavjud emas.")
                            formula_sheet = sheet_names[index]
                        _formula(formula if formula.startswith("=") else "=" + formula, bounds, formula_sheet)
                except ValueError as exc:
                    raise ValueError(
                        "Excel → PDF: xavfsiz qayta hisoblash uchun formula qo‘llab-quvvatlanmaydi. "
                        "Faqat fayl ichidagi oddiy formulalardan foydalaning. " + str(exc)
                    ) from exc


def _office_to_pdf(source: Path, destination: Path) -> str:
    if os.name != "nt":
        raise ValueError("DOCX/XLSX/PPTX → PDF uchun Windows va o‘rnatilgan Microsoft Office kerak.")
    _check_office_active_content(source)
    with tempfile.TemporaryDirectory(prefix="jarvis-pdf-", dir=destination.parent) as temporary:
        exported = Path(temporary) / "export.pdf"
        command = [sys.executable, "-m", "jarvis.office.convert_worker", str(source), str(exported)]
        try:
            result = subprocess.run(
                command, capture_output=True, text=True, encoding="utf-8", errors="replace",
                timeout=OFFICE_TIMEOUT_SECONDS, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except subprocess.TimeoutExpired as exc:
            raise ValueError("Microsoft Office PDF yaratish vaqt chegarasidan oshdi. Office aktivatsiyasi va faylni tekshiring.") from exc
        try:
            response = json.loads(result.stdout.strip())
        except (ValueError, TypeError):
            response = {}
        if result.returncode != 0 or not response.get("ok"):
            detail = str(response.get("error", "Microsoft Office bilan aloqa o‘rnatilmadi; Office va pywin32 o‘rnatilganini tekshiring."))[:700]
            raise ValueError(detail)
        if not exported.is_file() or exported.stat().st_size < 5:
            raise ValueError("Microsoft Office PDF faylini yaratmadi.")
        with exported.open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                raise ValueError("Microsoft Office natijasi haqiqiy PDF emas.")
            stream.seek(0)
            with destination.open("xb") as output:
                shutil.copyfileobj(stream, output)
    return str(response.get("backend", "Microsoft Office PDF"))


def convert_file(source: Path, destination: Path, *, sheet: str | None = None) -> str:
    """Convert to a NEW file only. This never replaces the source or an existing result."""
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if source == destination:
        raise ValueError("Manba fayl ustiga yozish mumkin emas; boshqa natija nomini tanlang.")
    if destination.exists():
        raise ValueError("Natija fayli allaqachon mavjud; boshqa nomni tanlang.")
    if not destination.parent.is_dir():
        raise ValueError("Natija uchun papka mavjud emas.")
    _check_source(source)
    extension, target = source.suffix.lower(), destination.suffix.lower()
    if target not in _CONVERSIONS.get(extension, ()):
        raise ValueError(f"{extension or '(formatsiz)'} → {target or '(formatsiz)'} aylantirish qo‘llab-quvvatlanmaydi.")
    if target == ".pdf":
        return _office_to_pdf(source, destination)
    if extension == ".pptx" and target == ".docx":
        return _pptx_to_docx(source, destination)
    if extension == ".csv":
        return _csv_to_xlsx(source, destination)
    if extension == ".xlsx":
        return _xlsx_to_csv(source, destination, sheet)
    if extension == ".txt":
        from docx import Document

        text = _read_text(source)
        document = Document()
        for line in text.split("\n"):
            document.add_paragraph(line.rstrip("\r"))
        with destination.open("xb") as stream:
            document.save(stream)
        return "python-docx (TXT → DOCX)"
    extractors = {".docx": (_docx_to_text, "python-docx"), ".pptx": (_pptx_to_text, "python-pptx"), ".pdf": (_pdf_to_text, "pypdf; faqat matn, OCR emas")}
    extractor, backend = extractors[extension]
    _write_text(destination, extractor(source))
    return f"{backend} ({extension.upper().lstrip('.')} → TXT; sahifa bezagi saqlanmaydi)"
